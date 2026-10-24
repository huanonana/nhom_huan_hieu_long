"""Apply final pagination to the report's manual contents and figure list."""
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
report = root / "output" / "report" / "Bao_cao_do_an_HHL_Stay.docx"
doc = Document(report)
for p in list(doc.paragraphs):
    text = p.text.strip()
    if text.startswith(("Phụ lục H: Firebase", "Phụ lục I: Kiểm tra Medium Phone", "Phụ lục J: Wireframe")) or (text.startswith("Hình J.") and "Wireframe" in text):
        p._element.getparent().remove(p._element)

for p in doc.paragraphs:
    if p.text.strip() == "Danh mục hình và bảng":
        for text in (
            "Phụ lục H: Firebase và bảo vệ dữ liệu\t64–69",
            "Phụ lục I: Kiểm tra Medium Phone\t70–71",
            "Phụ lục J: Wireframe sáu màn hình\t72–77",
        ):
            p.insert_paragraph_before(text)
        break

for p in doc.paragraphs:
    if p.text.startswith("Bảng F.1 User Stories"):
        for number, title in enumerate(("Home", "Explore", "Chi tiết", "Đặt phòng", "Đăng nhập", "Dashboard quản trị"), 1):
            p.insert_paragraph_before(f"Hình J.{number} Wireframe {title}\t{71 + number}")
        break
doc.save(report)
print(report)
