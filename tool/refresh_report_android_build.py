"""Update the existing final report with the verified Android build result."""
from pathlib import Path
from docx import Document

report = (
    Path(__file__).resolve().parents[1]
    / "output"
    / "report"
    / "Bao_cao_do_an_HHL_Stay.docx"
)
doc = Document(report)

for paragraph in doc.paragraphs:
    if "Build web Firebase-enabled đã hoàn tất" in paragraph.text:
        paragraph.text = paragraph.text.replace(
            "để build kiểm tra dùng `flutter build web --dart-define=HHL_USE_FIREBASE=true`. Build web Firebase-enabled đã hoàn tất thành công ngày 24/10/2026.",
            "để build kiểm tra dùng `flutter build web --dart-define=HHL_USE_FIREBASE=true` và `flutter build apk --debug --dart-define=HHL_USE_FIREBASE=true`. Cả Web và Android debug APK Firebase-enabled đều build thành công ngày 24/10/2026.",
        )

table = doc.tables[12]
command = "flutter build apk --debug --dart-define=HHL_USE_FIREBASE=true"
if not any(command in row.cells[0].text for row in table.rows):
    cells = table.add_row().cells
    cells[0].text = command
    cells[1].text = "Built app-debug.apk"
    cells[2].text = "Firebase-enabled Android debug compile/package; not installed on a device"
    for cell in cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = __import__("docx").shared.Pt(8)

doc.save(report)
print(report)
