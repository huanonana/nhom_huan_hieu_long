"""Refresh the Phenikaa report with verified Firebase/Medium Phone results and wireframes."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "output" / "report" / "Bao_cao_do_an_HHL_Stay.docx"
PHONE = ROOT / "docs" / "report_assets" / "15-medium-phone-firebase-explore.png"
WIREFRAMES = [
    ("01-home.png", "Home", "Trang chủ giới thiệu điểm đến và lối vào khám phá."),
    ("02-explore.png", "Explore", "Tìm kiếm, bộ lọc nhanh và danh sách homestay."),
    ("03-detail.png", "Chi tiết homestay", "Thông tin chỗ ở, tiện nghi, phòng và giá."),
    ("04-booking.png", "Đặt phòng", "Ngày ở, khách, mã ưu đãi và bước xác nhận giá."),
    ("05-login.png", "Đăng nhập", "Form đăng nhập và lối vào Host/Admin."),
    ("06-admin-dashboard.png", "Dashboard quản trị", "Số liệu vận hành và thao tác nhanh."),
]


def set_cell(row, first, values):
    for cell in row.cells:
        if cell.text.strip() == first:
            for p in cell.paragraphs:
                p._element.getparent().remove(p._element)
            for value in values:
                p = cell.add_paragraph(value)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
            return True
    return False


def style_picture_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)


doc = Document(REPORT)
paras = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}

replacements = {
    "BÁO CÁO ĐỒ ÁN CƠ SỞ\nĐỀ TÀI: XÂY DỰNG ỨNG DỤNG QUẢN LÝ HOMESTAY HHL STAY":
    "BÁO CÁO BÀI TẬP LỚN\nĐỀ TÀI: XÂY DỰNG ỨNG DỤNG QUẢN LÝ HOMESTAY HHL STAY",
    "Báo cáo gồm sáu chương và bảy phụ lục. Số trang được đối chiếu bản Word xuất PDF cuối, tính theo trang in. Các ảnh ghi lại giao diện app local; sơ đồ mô tả luồng hiện có và chỉ rõ giới hạn prototype.":
    "Báo cáo gồm sáu chương và mười phụ lục. Số trang được đối chiếu bản Word xuất PDF cuối, tính theo trang in. Báo cáo có ảnh app local/Medium Phone, UML, wireframe sáu màn hình và chỉ rõ giới hạn prototype.",
    "Phát triển ứng dụng web Flutter; tổ chức giao diện, trạng thái và repository; tạo danh mục mẫu; hỗ trợ tìm kiếm, yêu thích, gửi và xử lý đơn demo; trình bày màn vận hành.":
    "Phát triển ứng dụng Flutter Web/Android; tổ chức giao diện, trạng thái và repository; tạo danh mục mẫu; hỗ trợ tìm kiếm, yêu thích, gửi và xử lý đơn demo; trình bày màn vận hành.",
    "File firestore.rules trong repository giữ nguyên mô hình phát triển mà nhóm đã chọn: `allow read: if true` và `allow write: if request.auth != null` cho mọi document. Firebase Rules là điều kiện cấp quyền ở backend; UI ẩn nút hoặc router guard không thay thế kiểm tra ở server. Firebase mô tả rules là nơi kiểm tra trạng thái xác thực và dữ liệu trước khi cho phép đọc/ghi; xem tài liệu chính thức tại https://firebase.google.com/docs/firestore/security/rules-conditions.":
    "File firestore.rules giữ quy tắc demo: `allow read: if true` và `allow write: if request.auth != null` cho mọi document. Ngày 24/10/2026 nhóm triển khai rules lên project hhl-stay-home-2026-23010895 để kiểm tra client. Firebase Rules là điều kiện cấp quyền ở backend; UI ẩn nút hoặc router guard không thay thế kiểm tra server. Tham khảo https://firebase.google.com/docs/firestore/security/rules-conditions.",
    "Các rules này cho phép mọi người đọc toàn bộ dữ liệu, mọi tài khoản đã đăng nhập sửa/xóa mọi collection. Không được dùng cho dịch vụ công khai hoặc dữ liệu khách thật. File chỉ được chuẩn bị cục bộ; Codex không triển khai lên Firebase Console. Người quản lý phải xem nội dung, giới hạn quyền theo UID, role, chủ sở hữu và trạng thái, rồi tự kiểm tra bằng Rules Playground/Emulator trước khi phát hành.":
    "Rules này cho phép mọi người đọc toàn bộ dữ liệu, mọi tài khoản đã đăng nhập sửa/xóa mọi collection. Không dùng cho dịch vụ công khai hoặc dữ liệu khách thật. Kiểm tra ngày 24/10/2026: Firebase Auth tạo tài khoản thử, Firestore tạo/đọc hồ sơ thành công; tài khoản tạm được gửi lệnh xóa sau kiểm tra. Trước khi dùng dữ liệu thật, phải giới hạn quyền theo UID, role, chủ homestay và trạng thái.",
    "Trước khi đánh giá thật cần bổ sung kiểm thử Firebase Emulator cho tạo hồ sơ, booking được ghi/đọc sau đăng nhập, chuyển trạng thái, từ chối quyền và hai giao dịch đặt trùng. Bộ `flutter test` hiện chạy cô lập, không gửi request thật đến Firebase.":
    "Bộ `flutter test` chạy cô lập và không gửi request thật đến Firebase. Ngoài test tự động, ngày 24/10/2026 nhóm đã thử live Auth sign-up và Firestore profile create/read, rồi gửi yêu cầu dọn tài khoản tạm; việc xóa chưa được xác minh độc lập. Medium Phone Firebase mở được landing và Explore sau khi cập nhật Rules. Còn cần emulator/concurrency tests cho booking, chuyển trạng thái và từ chối quyền.",
    "Firebase network behavior is not covered by `flutter test`; tests run without Firebase initialization. The current build confirms Firebase-enabled web compilation, not successful live Auth/Firestore requests against every rule.":
    "Firebase network behavior is not covered by `flutter test`; tests run without Firebase initialization. Firebase Auth sign-up/profile create/read smoke passed and the Firebase-enabled Medium Phone loaded Explore without permission/timeout errors after Rules deployment. This does not cover concurrency or production authorization.",
    "Firebase-enabled Android debug compile/package; launcher smoke passed, live Firebase not tested":
    "Firebase-enabled Android debug build; live Auth/Firestore smoke passed",
}

for paragraph_index, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text in replacements:
        p.text = replacements[text]

# Correct the member story count, rubric status, and current integration evidence.
assert set_cell(doc.tables[10].rows[2], "Firebase Auth", ["Đăng ký/đăng nhập email-password; hồ sơ users", "Auth + hồ sơ thử live: pass"])
assert set_cell(doc.tables[10].rows[3], "Firestore booking", ["Tải/ghi booking; smoke profile Firestore đã pass", "Rules demo rộng; chưa kiểm thử concurrency"])
assert set_cell(doc.tables[10].rows[4], "Các document khác", ["Homestay, room, review, finance, staff, shift có mapper/lệnh ghi", "Cache local; chưa có queue/conflict resolver"])

table12 = doc.tables[12]
for row in table12.rows:
    if row.cells[0].text.strip() == "flutter build web --dart-define=HHL_USE_FIREBASE=true":
        row.cells[1].text = "Built build/firebase_web"
        row.cells[2].text = "Firebase-enabled Web compile/link"
    elif row.cells[0].text.strip() == "flutter build apk --debug --dart-define=HHL_USE_FIREBASE=true":
        row.cells[1].text = "Built; installed on emulator"
        row.cells[2].text = "Auth/Firestore smoke and Explore passed"
if not any(row.cells[0].text.strip() == "firebase deploy --only firestore:rules" for row in table12.rows):
    table12.add_row().cells[0].text = "firebase deploy --only firestore:rules"
    table12.rows[-1].cells[1].text = "Rules compiled and released"
    table12.rows[-1].cells[2].text = "Project lớp; rules demo only"

table13 = doc.tables[13]
for row in table13.rows:
    key = row.cells[0].text.strip()
    if key == "1. User Stories":
        row.cells[1].text = "19 stories (12 Guest + 7 Host)"
        row.cells[2].text = "Đủ tối thiểu nhóm 3/4 người; nhóm xác nhận phân công"
    elif key == "4. Screens/Flow":
        row.cells[1].text = "14 ảnh app, 6 wireframe và 7+ luồng"
        row.cells[2].text = "Xác nhận đúng mẫu Canvas nhóm chọn"
    elif key.startswith("6–8."):
        row.cells[1].text = "Flutter, Firebase Auth/Firestore, 96 test, Web + Android build"
        row.cells[2].text = "Còn emulator/concurrency và kiểm thử production"
    elif key.startswith("9–10."):
        row.cells[2].text = "Điền thông tin thật, link video/repo, đóng góp và ký nộp"

table14 = doc.tables[14]
for row in table14.rows:
    if row.cells[0].text.strip() == "flutter build apk --debug":
        row.cells[1].text = "Built local APK; Firebase APK built separately"
    elif row.cells[0].text.strip() == "Android route smoke":
        row.cells[1].text = "Firebase landing → Explore; no permission/timeout/crash logs"
        row.cells[2].text = "Auth/profile read/write smoke; booking path covered by widget tests"
    elif row.cells[0].text.strip() == "localhost 8083 / setup 8099":
        row.cells[2].text = "HTTP 200; local Web and setup service"
if not any(row.cells[0].text.strip() == "Firebase Auth + Firestore profile" for row in table14.rows):
    table14.add_row().cells[0].text = "Firebase Auth + Firestore profile"
    table14.rows[-1].cells[1].text = "Temporary account signup; profile create/read succeeded; cleanup requested"
    table14.rows[-1].cells[2].text = "No customer data; rules are classroom-wide"

# Update the risk matrix to distinguish local storage from Firebase mode.
for row in doc.tables[9].rows:
    if row.cells[0].text.strip() == "Tài khoản":
        row.cells[1].text = "Local demo hoặc Firebase Auth; profile Firebase ở users/{uid}."
        row.cells[2].text = "Rules hiện cho phép mọi người đọc và mọi account ghi."
        row.cells[3].text = "Rules theo UID/role/chủ sở hữu; audit và recovery."
    elif row.cells[0].text.strip() == "Lưu trữ":
        row.cells[1].text = "SharedPreferences local cache; Firebase mode lưu các document đã hỗ trợ."
        row.cells[2].text = "Chưa có retry queue hay conflict resolution đa thiết bị."

# Replace the old local-only evidence paragraph and screenshot in Appendix I.
for paragraph_index, p in enumerate(doc.paragraphs):
    if p.text.startswith("Bản local được build lại từ workspace"):
        p.text = (
            "Bản Web local được build và đang phục vụ tại http://127.0.0.1:8083; "
            "Firebase Web được build riêng vào build/firebase_web. Hai Android debug APK "
            "(local/Firebase) được build trên cùng source. Medium Phone emulator chạy "
            "Firebase build: landing → Explore, không còn banner sync; sau cập nhật Rules "
            "logcat không có PERMISSION_DENIED, timeout hay crash. Auth sign-up và Firestore "
            "profile create/read smoke đã qua."
        )
    if p.text.startswith("Hình I.1."):
        p.text = "Hình I.1. Explore trên Medium Phone với Firebase mode sau khi Rules được áp dụng."
        current_paragraphs = doc.paragraphs
        if paragraph_index > 0 and current_paragraphs[paragraph_index - 1]._element.xpath(".//w:drawing"):
            old_picture = current_paragraphs[paragraph_index - 1]
            old_picture._element.getparent().remove(old_picture._element)
        image = p.insert_paragraph_before()
        style_picture_paragraph(image)
        image.add_run().add_picture(str(PHONE), width=Inches(1.9))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Refresh contents and add appendix references before the figure/table index.
for p in doc.paragraphs:
    if p.text.strip() in replacements:
        p.text = replacements[p.text.strip()]
    if p.text.startswith("Các hình UML được xây dựng"):
        p.text = ("Các hình UML đối chiếu nghiệp vụ với model/repository. Báo cáo có một Class Diagram, "
                  "một Use Case Diagram, ba Activity, bốn Sequence, một vòng đời booking và sáu wireframe mobile. "
                  "Các hình được liệt kê dưới đây.")

for p in doc.paragraphs:
    if p.text.startswith("Phụ lục G: Ảnh chụp giao diện"):
        p.text = "Phụ lục G: Ảnh chụp giao diện\t57–63"
    if p.text.strip() == "Danh mục hình và bảng":
        if not any(q.text.startswith("Phụ lục H: Firebase") for q in doc.paragraphs):
            p.insert_paragraph_before("Phụ lục H: Firebase và bảo vệ dữ liệu\t64–69")
            p.insert_paragraph_before("Phụ lục I: Kiểm tra Medium Phone\t70–71")
            p.insert_paragraph_before("Phụ lục J: Wireframe sáu màn hình\t72–77")
    if p.text.startswith("Bảng F.1 User Stories"):
        if not any(q.text.startswith("Hình J.1 Wireframe") for q in doc.paragraphs):
            for i, title in enumerate(("Home", "Explore", "Chi tiết", "Đặt phòng", "Đăng nhập", "Dashboard quản trị"), 1):
                p.insert_paragraph_before(f"Hình J.{i} Wireframe {title}\t{71+i}")

# Appendix J: one wireframe per page for legibility and print review.
if not any(p.text.strip() == "PHỤ LỤC J. WIREFRAME SÁU MÀN HÌNH" for p in doc.paragraphs):
    doc.add_page_break()
    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("PHỤ LỤC J. WIREFRAME SÁU MÀN HÌNH")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(48, 54, 125)
    intro = doc.add_paragraph(
        "Wireframe mobile do nhóm phác thảo cho HHL Stay. Khung ảnh là vị trí chờ ảnh homestay trong bản bố cục; "
        "nội dung và ảnh thật của ứng dụng được trình bày ở Phụ lục G và I. Đây là mockup của nhóm, không thay cho việc xác nhận mẫu Canvas giảng viên đã chỉ định."
    )
    intro.paragraph_format.space_after = Pt(6)
    for run in intro.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
    for index, (filename, name, description) in enumerate(WIREFRAMES):
        if index:
            doc.add_page_break()
        heading = doc.add_paragraph()
        heading.style = "Heading 2"
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run(f"J.{index + 1}. {name}")
        detail = doc.add_paragraph(description)
        detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in detail.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(90, 105, 130)
        image = doc.add_paragraph()
        style_picture_paragraph(image)
        image.add_run().add_picture(str(ROOT / "docs" / "wireframes" / filename), width=Inches(3.25))
        caption = doc.add_paragraph(f"Hình J.{index + 1}. Wireframe {name} — HHL Stay.")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.italic = True

doc.save(REPORT)
print(f"Updated report: {REPORT}")
