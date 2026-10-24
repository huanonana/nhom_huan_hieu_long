"""Update the existing capstone report with the Firebase implementation addendum."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "output" / "report" / "Bao_cao_do_an_HHL_Stay.docx"


def replace_in_paragraph(paragraph, replacements):
    original = paragraph.text
    updated = original
    for old, new in replacements:
        updated = updated.replace(old, new)
    if updated != original:
        paragraph.text = updated


def style_cell(cell, fill=None, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        tc_pr = cell._tc.get_or_add_tcPr()
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), fill)
        tc_pr.append(shade)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(9)


def page_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = RGBColor(48, 54, 125)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.style = "Subtitle"
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(90, 105, 130)


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
        style_cell(t.rows[0].cells[i], "E8EAFB", True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            style_cell(cells[i])
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return t


def add_page(doc, title, subtitle=None):
    doc.add_page_break()
    page_title(doc, title, subtitle)


doc = Document(REPORT)

# Bring the report's status statements forward to the verified 24 October build.
replacements = [
    ("23/10/2026", "24/10/2026"),
    (
        "Homestay, giá, tài khoản và booking mẫu là dữ liệu minh họa. Các giới hạn chưa kiểm chứng được ghi rõ, không khẳng định quá khả năng của bản thử.",
        "Dữ liệu homestay và booking mẫu là dữ liệu minh họa. Chế độ local và chế độ Firebase opt-in được phân biệt rõ; phần chưa kiểm chứng live được ghi riêng.",
    ),
    (
        "HHL Stay là ứng dụng web thử nghiệm hỗ trợ khám phá homestay, xem phòng, gửi booking và thao tác quản trị cơ bản. Ứng dụng dùng Flutter/Dart, Riverpod, go_router và SharedPreferences. Báo cáo trình bày yêu cầu, mô hình, giao diện, triển khai và kiểm thử. Phạm vi hiện tại là prototype cục bộ, chưa có backend, xác thực máy chủ, thanh toán hoặc đồng bộ đa thiết bị.",
        "HHL Stay là ứng dụng Flutter hỗ trợ khám phá homestay, gửi booking và thao tác quản trị. Ứng dụng dùng Riverpod và go_router; local mode lưu SharedPreferences, Firebase mode opt-in dùng Firebase Authentication và Firestore cho hồ sơ, đơn đặt và các thay đổi domain được hỗ trợ. Chưa có thanh toán thật, xử lý đặt phòng nguyên tử phía máy chủ hoặc bộ giải quyết xung đột offline.",
    ),
    (
        "HHL Stay is a Flutter Web prototype for discovering homestays, reviewing rooms, submitting booking requests, and trying basic host operations. It uses Riverpod, go_router, and device-local SharedPreferences storage. This report covers user stories, requirements, UML, implementation, and tests. The current build is an academic demo and does not provide production authentication, shared backend, or payment processing.",
        "HHL Stay is a Flutter app for discovering homestays, submitting bookings, and host operations. Local mode uses SharedPreferences; an opt-in Firebase mode uses Firebase Authentication and Firestore for profiles, booking documents, and supported domain changes. The academic demo has no real payment processing, server-side atomic booking transaction, or production-grade security rules.",
    ),
    (
        "Tài liệu hiện hành nói phiên demo cục bộ, không Firebase Auth/server auth. Không thể khẳng định admin được bảo vệ trước sửa client hoặc dữ liệu an toàn. Không tạo tài khoản thật khi lập báo cáo.",
        "App có Firebase Auth/Firestore trong chế độ bật bằng `--dart-define=HHL_USE_FIREBASE=true`; local mode vẫn dùng tài khoản demo. Chưa kiểm thử network auth/rules bằng Firebase Emulator trong bộ test tự động. Rules dev cho phép đọc công khai và mọi tài khoản đăng nhập ghi dữ liệu, nên không an toàn cho sản phẩm thật.",
    ),
    (
        "Thêm Auth/Firestore hoặc API, rules, transaction booking, image storage và emulator tests. Đưa quyền, giá, QR, tiền nhạy cảm lên server; thử hai yêu cầu đặt trùng đồng thời.",
        "Tiếp theo cần thay rules rộng bằng phân quyền theo UID/chủ homestay, thêm transaction chống đặt trùng, kiểm thử Firebase Emulator, lưu ảnh an toàn, và kiểm tra role/giá/QR/tài chính ở server.",
    ),
    (
        "Từ điển dưới đây lấy từ serialization model hiện tại. Kiểu được ghi ở mức khái niệm; JSON được lưu ở SharedPreferences, chưa phải schema Firestore hay hợp đồng API.",
        "Từ điển dưới đây lấy từ serialization model hiện tại. JSON tiếp tục dùng trong SharedPreferences local mode và được ánh xạ thành document Firestore trong Firebase mode; chưa phải hợp đồng API production.",
    ),
    (
        "Workspace 24/10/2026: analyzer không lỗi, 96 test pass. Test cô lập không kiểm backend hay đồng thời nhiều thiết bị.",
        "Workspace 24/10/2026: `flutter analyze` không lỗi, 96 test pass và `flutter build web --dart-define=HHL_USE_FIREBASE=true` thành công. Test cô lập không kiểm network Firebase, rules thật hay đồng thời nhiều thiết bị.",
    ),
]
for p in doc.paragraphs:
    replace_in_paragraph(p, replacements)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_in_paragraph(p, replacements)

add_page(doc, "PHỤ LỤC H. CẬP NHẬT FIREBASE", "Bản cập nhật kỹ thuật — 24/10/2026")
body(
    doc,
    "Bản báo cáo chính được lập khi app chỉ chạy local. Phụ lục H cập nhật đúng phần Firebase vừa được tích hợp và có hiệu lực khi bật HHL_USE_FIREBASE=true. Nếu nội dung trước đó nói app chưa có Firebase, hiểu là mô tả trạng thái cũ của bản local; các giới hạn bảo mật và vận hành nêu ở đây là trạng thái hiện hành.",
)
body(
    doc,
    "Firebase project đã cấu hình cho Android và Web: hhl-stay-home-2026-23010895. Firebase Core được khởi tạo từ firebase_options.dart. FirebaseAuthRepository tạo tài khoản email/mật khẩu và hồ sơ users/{uid}; mật khẩu do Firebase Authentication quản lý, không được serialize vào hồ sơ Firestore. Tài khoản local demo không tự chuyển thành tài khoản Firebase.",
)
table(
    doc,
    ["Phần", "Trạng thái hiện thực", "Giới hạn"],
    [
        ("Local mode", "Mặc định; SharedPreferences và tài khoản demo", "Không đồng bộ nhiều thiết bị"),
        ("Firebase Auth", "Đăng ký/đăng nhập email-password; hồ sơ users", "Chưa test emulator/network"),
        ("Firestore booking", "Tải booking lúc khởi động; ghi đơn và cập nhật trạng thái", "Tải toàn bộ collection theo rules dev"),
        ("Các document khác", "Homestay, room, review, finance, staff, shift có mapper và lệnh ghi", "Không có queue/conflict resolver"),
    ],
    [1.35, 3.05, 2.2],
)

add_page(doc, "H.1. CẤU HÌNH VÀ LUỒNG TÀI KHOẢN")
body(doc, "Chạy local như trước bằng `flutter run -d web-server --web-hostname 127.0.0.1 --web-port 8083`. Để bật Firebase trên Chrome dùng `flutter run -d chrome --dart-define=HHL_USE_FIREBASE=true`; để build kiểm tra dùng `flutter build web --dart-define=HHL_USE_FIREBASE=true`. Build web Firebase-enabled đã hoàn tất thành công ngày 24/10/2026.")
body(doc, "Khi đăng ký, Firebase Authentication cấp UID. App lưu uid, email, displayName, role và createdAt trong users/{uid}. Khi đăng nhập, stream authStateChanges đọc hồ sơ tương ứng; nếu hồ sơ thiếu, app tạo hồ sơ mặc định guest. Đăng xuất chuyển phiên Firebase về null và router bảo vệ quay về luồng khách chưa xác thực.")
table(
    doc,
    ["Collection", "Document ID / trường chính", "Logic dùng"],
    [
        ("users", "UID; email, displayName, role, createdAt", "Hồ sơ đăng ký/đăng nhập"),
        ("bookings", "guestUID + mã đơn; các trường từ Booking.toJson", "Lưu đơn, đọc đơn khi khởi động và ghi chuyển trạng thái"),
        ("homestays / rooms", "ID domain ổn định; model JSON", "Mapper CRUD cho thao tác chủ nhà"),
        ("reviews", "bookingId; điểm, bình luận, guestId", "Ghi nhận đánh giá sau khi nghiệp vụ cho phép"),
        ("financeEntries / staffMembers / staffShifts", "ID domain ổn định", "Ghi các thay đổi mới từ màn quản trị"),
    ],
    [1.8, 2.4, 2.4],
)

add_page(doc, "H.2. ÁNH XẠ DOMAIN VÀ CƠ CHẾ LƯU")
body(doc, "FirebaseDocumentStore dùng các model domain hiện có thay vì đưa logic Flutter widget vào document. Mỗi model chuyển qua toJson; đọc Booking.fromJson, Homestay.fromJson, Room.fromJson và StayReview.fromJson. ID booking trên Firestore ghép guest UID với mã local nhằm tránh va chạm số thứ tự HHL-0001 được tạo độc lập trên nhiều thiết bị.")
body(doc, "BookingFormPage giữ bước xác nhận, tạo đơn local trong bộ nhớ, sau đó chờ Firestore ghi. Nếu Firestore từ chối, app khôi phục snapshot repository và không báo đặt thành công. Các chuyển trạng thái booking, ghi thanh toán, check-in/out, sửa homestay/phòng và đánh giá gọi đồng bộ nền; lỗi được đưa ra notice để người vận hành biết local đã đổi nhưng cloud chưa nhận.")
body(doc, "DemoStorage vẫn giữ cache local và các danh sách yêu thích/lịch sử xem. Khi mở Firebase mode, app tải booking cloud hiện có rồi ghép theo ID, giữ lại dữ liệu demo chưa có trên cloud. Đây chưa phải đồng bộ hai chiều hoàn chỉnh: catalog công khai chưa tải về để thay thế fixtures, không có retry queue khi offline, xung đột giữa hai thiết bị chưa được giải quyết, và tài chính/staff chưa có truy vấn phân trang theo tenant.")
body(doc, "Trước khi đánh giá thật cần bổ sung kiểm thử Firebase Emulator cho tạo hồ sơ, booking được ghi/đọc sau đăng nhập, chuyển trạng thái, từ chối quyền và hai giao dịch đặt trùng. Bộ `flutter test` hiện chạy cô lập, không gửi request thật đến Firebase.")

add_page(doc, "H.3. RULES VÀ BẢO VỆ DỮ LIỆU")
body(doc, "File firestore.rules trong repository giữ nguyên mô hình phát triển mà nhóm đã chọn: `allow read: if true` và `allow write: if request.auth != null` cho mọi document. Firebase Rules là điều kiện cấp quyền ở backend; UI ẩn nút hoặc router guard không thay thế kiểm tra ở server. Firebase mô tả rules là nơi kiểm tra trạng thái xác thực và dữ liệu trước khi cho phép đọc/ghi; xem tài liệu chính thức tại https://firebase.google.com/docs/firestore/security/rules-conditions.")
body(doc, "Các rules này cho phép mọi người đọc toàn bộ dữ liệu, mọi tài khoản đã đăng nhập sửa/xóa mọi collection. Không được dùng cho dịch vụ công khai hoặc dữ liệu khách thật. File chỉ được chuẩn bị cục bộ; Codex không triển khai lên Firebase Console. Người quản lý phải xem nội dung, giới hạn quyền theo UID, role, chủ sở hữu và trạng thái, rồi tự kiểm tra bằng Rules Playground/Emulator trước khi phát hành.")
body(doc, "Không đưa service account key vào Flutter hoặc repository. Firebase client config trong app không phải khóa quyền quản trị; API key Gemini cũng không được coi là bí mật nếu nhúng vào Web. Chỉ dùng dữ liệu giả trong demo, không thu thập giấy tờ/địa chỉ riêng tư không cần thiết, cho phép xóa tài khoản và dữ liệu, và nêu mục đích lưu dữ liệu cho người dùng.")
body(doc, "Đề bài nhắc một tin Facebook/Cambridge Analytica “22/10/2026, 25 tỷ USD”. Nguồn FTC chính thức được kiểm tra ghi nhận Facebook chấp nhận khoản phạt 5 tỷ USD trong dàn xếp năm 2019; FTC cũng có hồ sơ riêng về Cambridge Analytica và việc thu thập dữ liệu người dùng. Báo cáo không lặp lại con số/nội dung 2026 khi chưa có nguồn xác minh. Tham khảo: https://www.ftc.gov/news-events/news/press-releases/2019/07/ftc-imposes-5-billion-penalty-sweeping-new-privacy-restrictions-facebook và https://www.ftc.gov/news-events/news/press-releases/2019/07/ftc-sues-cambridge-analytica-settles-former-ceo-app-developer.")

add_page(doc, "H.4. KIỂM THỬ VÀ TÁI LẬP BUILD")
table(
    doc,
    ["Lệnh", "Kết quả 24/10/2026", "Phạm vi chứng minh"],
    [
        ("flutter analyze", "No issues found", "Static analysis trên workspace"),
        ("flutter test", "All tests passed — 96", "Unit/widget flows local"),
        ("flutter build web --dart-define=HHL_USE_FIREBASE=true", "Built build/web", "Firebase-enabled web compile/link"),
    ],
    [2.45, 1.65, 2.5],
)
body(doc, "Các kiểm thử hiện có bao phủ điều hướng, đăng nhập demo, khám phá, bộ lọc, chi tiết, tính ngày và giá, booking, trạng thái đơn, review, thao tác host, lịch, tài chính, AI fallback và bố cục hẹp. Firebase mode build được nhưng chưa có kiểm thử tích hợp cho mạng/auth/rules; vì vậy không diễn giải kết quả compile như bằng chứng đã ghi dữ liệu thật thành công.")
body(doc, "Tài liệu Firebase Flutter chính thức hướng dẫn đăng ký app Android/Web, cấu hình FlutterFire và khởi tạo Firebase: https://firebase.google.com/docs/flutter/setup. Để tái lập, chạy `flutter pub get`, chọn đúng project Firebase đã cấp, bật Email/Password trong Authentication, kiểm tra Firestore database và rules, rồi chạy build flag nêu tại H.1. Không đưa key bí mật hoặc thông tin đăng nhập cá nhân vào log/báo cáo.")

add_page(doc, "H.5. ĐỐI CHIẾU ĐỀ VÀ VIỆC CẦN XÁC NHẬN")
table(
    doc,
    ["Mục đề", "Bằng chứng trong repo/báo cáo", "Còn phải xác nhận"],
    [
        ("1. User Stories", "12 story và tiêu chí nghiệm thu", "Nếu nhóm 4 người, tăng tối thiểu 16"),
        ("2–3. Phân tích, UML", "Đối tượng, quan hệ, class diagram, 7 sơ đồ thuật toán", "Kiểm tra cách vẽ/diễn giải với giảng viên"),
        ("4. Screens/Flow", "Ảnh 14 màn hình và 5+ luồng", "Thêm Figma nếu cần file editable"),
        ("5. Layout mẫu", "App screenshot và mockup theo HHL Stay", "Ảnh Canvas nguồn không mở được; xác nhận màn hình mẫu được chọn"),
        ("6–8. Code, Firebase, Test", "Flutter, Firebase mode, 96 test, web build", "Chạy thử auth/Firestore Emulator và mobile thật"),
        ("9–10. Nộp và đóng góp", "Báo cáo có bìa, demo flow, bảng thành viên", "Xác nhận tên/MSSV, tỷ lệ đóng góp, GitHub/demo link và ký nộp"),
    ],
    [1.35, 2.75, 2.5],
)
body(doc, "Bìa và bảng thành viên đang giữ thông tin đọc được từ repository nhưng phải được từng thành viên kiểm tra; tỷ lệ đóng góp chưa được xác nhận nên không tự bịa. Trước khi nộp PDF/bản cứng, cập nhật link demo dưới 8 phút, link GitHub, giảng viên, tên và MSSV từng bản cá nhân. Bản cứng/PDF theo đề: trước 17:00 và 23:59 ngày 03/11/2026.")

REPORT.parent.mkdir(parents=True, exist_ok=True)
doc.save(REPORT)
print(REPORT)
