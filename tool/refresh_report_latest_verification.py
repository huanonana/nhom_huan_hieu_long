"""Refresh the Phenikaa report with the latest verified local/Android results."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "output" / "report" / "Bao_cao_do_an_HHL_Stay.docx"
doc = Document(REPORT)

for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text.startswith("Nhóm đã xây dựng app Flutter cho khám phá"):
        paragraph.text = (
            "Nhóm đã xây dựng app Flutter cho khám phá, đặt phòng demo và quản trị vận hành. "
            "Lần kiểm tra 24/10/2026: flutter analyze sạch, 96/96 test tự động qua, "
            "Web local, Android local/Firebase debug APK và local release APK build thành công. "
            "Firebase debug mở được Explore trên Medium Phone; app process không ghi nhận Flutter "
            "fatal, Firestore permission-denied hay timeout. Khi thử Firebase release, máy chủ "
            "gần cạn RAM trong bước tối ưu và Android emulator báo System UI/Phone service ANR; "
            "màn hình HHL Stay vẫn được render. Firebase release APK chưa hoàn tất. Đây là giới hạn "
            "môi trường kiểm thử, không phải kết quả xác nhận production."
        )
    elif text.startswith("Bản Web local được build và đang phục vụ"):
        paragraph.text = (
            "Web local được build và đang phục vụ tại http://127.0.0.1:8083; hai Android debug APK "
            "(local/Firebase) và local release APK build thành công. Firebase debug đã mở landing "
            "và Explore trên Medium Phone; Android app log không có Flutter crash, permission-denied "
            "hay timeout. Ở lần thử Firebase release kế tiếp, host gần cạn RAM và emulator phát "
            "cảnh báo System UI/Phone service ANR trong khi HHL Stay vẫn hiện trên màn hình; Firebase "
            "release build đã dừng, chưa có APK release Firebase. Auth sign-up và Firestore profile "
            "create/read smoke đã qua; chưa kiểm thử concurrency."
        )

table13 = doc.tables[13]
for row in table13.rows:
    key = row.cells[0].text.strip()
    if key.startswith("6–8."):
        row.cells[1].text = (
            "Flutter, Firebase Auth/Firestore, 96 test, Web, local/Firebase debug APK, local release"
        )
        row.cells[2].text = (
            "Firebase concurrency/production chưa test; emulator OS ANR khi host cạn RAM trong lần release thử"
        )

table14 = doc.tables[14]
for row in table14.rows:
    key = row.cells[0].text.strip()
    if key == "flutter build apk --debug":
        row.cells[1].text = "Local và Firebase debug APK build thành công"
        row.cells[2].text = "Debug artifacts trong output/builds; Firebase APK smoke-tested"
    elif key == "Android route smoke":
        row.cells[1].text = (
            "Firebase debug mở landing/Explore; local release render landing; app log không có Flutter fatal"
        )
        row.cells[2].text = (
            "Lần Firebase release thử gây Android System UI/Phone-service ANR khi host gần cạn RAM; "
            "HHL Stay vẫn render; Firebase release chưa hoàn tất"
        )
    elif key == "localhost 8083 / setup 8099":
        row.cells[1].text = "HTTP 200 / HTTP 200"
        row.cells[2].text = "Web local và trang setup trả mã 200 tại lần kiểm tra 24/10/2026"

if not any(row.cells[0].text.strip() == "flutter build apk --release (local)" for row in table14.rows):
    cells = table14.add_row().cells
    cells[0].text = "flutter build apk --release (local)"
    cells[1].text = "Built local release APK; app rendered on Medium Phone"
    cells[2].text = "65.2 MB; unsigned demo build, not a store release"

for table in (table13, table14):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

doc.save(REPORT)
print(REPORT)
